backend/scripts/import_legacy_stage.py#!/usr/bin/env python
"""Importa archivos XLSX legacy a tablas staging_* (carga cruda e idempotente)."""
from __future__ import annotations

import argparse
import logging
import os
import re
import sys
from datetime import datetime
from pathlib import Path
from typing import Iterable

import django
from django.db import connection, transaction

from legacy_migration_utils import IncidentLogger, normalize_string

try:
    import openpyxl
except ImportError as exc:
    raise SystemExit("Falta openpyxl. Instala con: pip install openpyxl") from exc


LOGGER = logging.getLogger(__name__)

EXPECTED_FILES = {
    "dbo_articulos.xlsx",
    "dbo_categorias.xlsx",
    "dbo_contactos.xlsx",
    "dbo_datosempresa.xlsx",
    "dbo_empleados.xlsx",
    "dbo_impuestos.xlsx",
    "dbo_ivas.xlsx",
    "dbo_ivas_r.xlsx",
    "dbo_motos_registradas.xlsx",
    "dbo_numeracion_fac.xlsx",
    "dbo_prefactura.xlsx",
    "dbo_usuarios.xlsx",
    "dbo_vendedores.xlsx",
}


def normalize_identifier(name: str) -> str:
    cleaned = re.sub(r"[^a-zA-Z0-9_]+", "_", name.strip())
    cleaned = re.sub(r"_+", "_", cleaned).strip("_")
    return cleaned.lower() or "col"


def table_name_for_file(path: Path) -> str:
    return f"staging_{normalize_identifier(path.stem)}"


def unique_columns(headers: Iterable[object]) -> list[str]:
    seen: dict[str, int] = {}
    columns: list[str] = []
    for idx, raw in enumerate(headers, start=1):
        base = normalize_identifier(normalize_string(raw) or f"col_{idx}")
        seen[base] = seen.get(base, 0) + 1
        col = base if seen[base] == 1 else f"{base}_{seen[base]}"
        columns.append(col)
    return columns


def list_xlsx_files(data_dir: Path) -> tuple[list[Path], list[Path]]:
    all_xlsx = sorted(path for path in data_dir.glob("*.xlsx"))
    ignored = [path for path in all_xlsx if path.name.startswith("~$")]
    valid = [path for path in all_xlsx if not path.name.startswith("~$")]
    return valid, ignored


def table_exists(table: str) -> bool:
    with connection.cursor() as cursor:
        cursor.execute(
            """
            SELECT EXISTS (
                SELECT 1 FROM information_schema.tables
                WHERE table_schema = 'public' AND table_name = %s
            )
            """,
            [table],
        )
        return bool(cursor.fetchone()[0])


def ensure_table(table: str, columns: list[str]) -> None:
    quote = connection.ops.quote_name
    with connection.cursor() as cursor:
        if not table_exists(table):
            cols_sql = ", ".join(f"{quote(col)} TEXT" for col in columns)
            cursor.execute(
                f"""
                CREATE TABLE {quote(table)} (
                    _source_file TEXT NOT NULL,
                    _source_row INTEGER NOT NULL,
                    _loaded_at TIMESTAMP NOT NULL,
                    {cols_sql}
                )
                """
            )
            LOGGER.info("Creada tabla staging %s", table)
            return

        cursor.execute(
            """
            SELECT column_name
            FROM information_schema.columns
            WHERE table_schema = 'public' AND table_name = %s
            ORDER BY ordinal_position
            """,
            [table],
        )
        existing = {r[0] for r in cursor.fetchall()}
        for col in columns:
            if col not in existing:
                cursor.execute(f"ALTER TABLE {quote(table)} ADD COLUMN {quote(col)} TEXT")
                LOGGER.info("Tabla %s: columna agregada %s", table, col)


def clear_table(table: str) -> int:
    quote = connection.ops.quote_name
    with connection.cursor() as cursor:
        cursor.execute(f"SELECT COUNT(*) FROM {quote(table)}")
        rows = int(cursor.fetchone()[0] or 0)
        if rows:
            cursor.execute(f"TRUNCATE TABLE {quote(table)}")
        return rows


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--data-dir",
        type=Path,
        default=Path(__file__).resolve().parents[2] / "data",
        help="Directorio con XLSX legacy",
    )
    parser.add_argument("--commit", action="store_true", help="Confirma inserción")
    parser.add_argument("--dry-run", action="store_true", help="Forzar rollback")
    parser.add_argument(
        "--incidents-json",
        type=Path,
        default=Path(__file__).resolve().parents[1] / "logs" / "legacy_stage_incidents.json",
    )
    return parser.parse_args()


def excel_rows(path: Path) -> tuple[list[str], list[list[str]]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    stream = ws.iter_rows(values_only=True)
    headers = next(stream, None)
    if not headers:
        return [], []
    columns = unique_columns(headers)
    rows_out: list[list[str]] = []
    for row in stream:
        normalized = [normalize_string(v) for v in (row or ())]
        if len(normalized) < len(columns):
            normalized.extend([""] * (len(columns) - len(normalized)))
        normalized = normalized[: len(columns)]
        if any(normalized):
            rows_out.append(normalized)
    return columns, rows_out


def insert_rows(table: str, source_file: str, columns: list[str], rows: list[list[str]]) -> int:
    if not rows:
        return 0
    quote = connection.ops.quote_name
    all_cols = ["_source_file", "_source_row", "_loaded_at", *columns]
    col_sql = ", ".join(quote(c) for c in all_cols)
    placeholders = ", ".join(["%s"] * len(all_cols))
    sql = f"INSERT INTO {quote(table)} ({col_sql}) VALUES ({placeholders})"

    payload = []
    now = datetime.utcnow()
    for i, row in enumerate(rows, start=2):
        payload.append([source_file, i, now, *row])

    with connection.cursor() as cursor:
        cursor.executemany(sql, payload)
    return len(payload)


def main() -> None:
    logging.basicConfig(level=logging.INFO, format="[%(levelname)s] %(message)s")
    args = parse_args()
    dry_run = (not args.commit) or args.dry_run

    backend_dir = Path(__file__).resolve().parents[1]
    if str(backend_dir) not in sys.path:
        sys.path.insert(0, str(backend_dir))
    os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings")
    django.setup()

    incidents = IncidentLogger()
    data_dir = args.data_dir
    if not data_dir.exists():
        raise SystemExit(f"No existe data-dir: {data_dir}")

    files, ignored = list_xlsx_files(data_dir)
    if not files and not ignored:
        raise SystemExit(f"No se encontraron .xlsx en {data_dir}")

    LOGGER.info("Data dir: %s", data_dir)
    LOGGER.info("Archivos encontrados (.xlsx): %s", len(files) + len(ignored))
    LOGGER.info("Archivos ignorados (temporales ~$): %s", len(ignored))
    for path in ignored:
        LOGGER.info("IGNORADO: %s", path.name)
    LOGGER.info("Archivos válidos a procesar: %s", len(files))
    for path in files:
        LOGGER.info("DETECTADO: %s -> %s", path.name, table_name_for_file(path))

    total_rows = 0
    processed = 0
    failed = 0

    with transaction.atomic():
        for path in files:
            try:
                columns, rows = excel_rows(path)
                if not columns:
                    incidents.add("WARN", "STAGE", path.name, "", "Archivo vacío o sin encabezados")
                    LOGGER.warning("Sin encabezados o vacío: %s", path.name)
                    processed += 1
                    continue

                table = table_name_for_file(path)
                ensure_table(table, columns)
                previous = clear_table(table)
                inserted = insert_rows(table, path.name, columns, rows)
                total_rows += inserted
                processed += 1
                LOGGER.info("PROCESADO: %s -> %s (previas=%s, insertadas=%s)", path.name, table, previous, inserted)
            except Exception as exc:  # noqa: BLE001 - necesitamos registrar y continuar con otros archivos
                failed += 1
                incidents.add("ERROR", "STAGE", path.name, "", f"Fallo procesando archivo: {exc}")
                LOGGER.exception("FALLIDO: %s", path.name)

        incidents.dump_json(args.incidents_json)

        LOGGER.info("=== Resumen staging ===")
        LOGGER.info("Total archivos detectados: %s", len(files) + len(ignored))
        LOGGER.info("Total archivos procesados: %s", processed)
        LOGGER.info("Total archivos ignorados: %s", len(ignored))
        LOGGER.info("Total archivos fallidos: %s", failed)
        LOGGER.info("Total filas insertadas staging: %s", total_rows)

        if dry_run:
            LOGGER.warning("Dry-run activo: rollback")
            transaction.set_rollback(True)


if __name__ == "__main__":
    main()
